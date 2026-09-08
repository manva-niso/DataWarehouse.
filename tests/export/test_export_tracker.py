from datetime import date, timedelta
from unittest.mock import patch

import pytest
from openpyxl import load_workbook

from export.export_tracker import (
    FutureDateError,
    PostingNotFoundError,
    InvalidStatusError,
    export_tracker_to_excel,
    mark_application,
)


def test_export_tracker_to_excel_zero_rows_writes_valid_headers(tmp_path):
    output = tmp_path / "tracker.xlsx"
    with patch("export.export_tracker.query_rows", return_value=[]):
        result = export_tracker_to_excel(str(output))
    assert result == str(output)
    workbook = load_workbook(str(output), read_only=True)
    sheet = workbook["Application Tracker"]
    headers = [cell.value for cell in next(sheet.iter_rows())]
    assert headers[0] == "Posting ID"
    assert "Closing Date" in headers
    assert len(list(sheet.iter_rows())) == 1


def test_export_tracker_to_excel_missing_closing_date_renders_not_specified(tmp_path):
    output = tmp_path / "tracker.xlsx"
    rows = [
        {
            "posting_id": "adzuna-1",
            "company_name": "Acme",
            "job_title": "Data Engineer",
            "location": "London",
            "date_posted": date(2026, 8, 1),
            "closing_date": None,
            "date_applied": None,
            "is_likely_closed": False,
            "last_seen_at": None,
        }
    ]
    with patch("export.export_tracker.query_rows", return_value=rows):
        export_tracker_to_excel(str(output))
    workbook = load_workbook(str(output), read_only=True)
    sheet = workbook["Application Tracker"]
    values = [cell.value for cell in next(sheet.iter_rows(min_row=2))]
    assert values[5] == "Not specified"
    assert values[6] in (None, "")


def test_export_tracker_to_excel_locked_file_retries_with_timestamp_suffix(tmp_path):
    original = tmp_path / "tracker.xlsx"
    stamped = tmp_path / "tracker-20260812120000.xlsx"
    with patch(
        "export.export_tracker._export_tracker_to_excel",
        side_effect=[PermissionError("locked"), str(stamped)],
    ):
        result = export_tracker_to_excel(str(original))
    assert result != str(original)
    assert "-20260812120000" in result


def test_mark_application_future_date_raises_clear_error():
    with pytest.raises(FutureDateError):
        mark_application("adzuna-1", date.today() + timedelta(days=1))


def test_mark_application_unknown_posting_raises_specific_error():
    with patch("export.export_tracker.query_rows", return_value=[]):
        with pytest.raises(PostingNotFoundError):
            mark_application("adzuna-999", date(2026, 8, 1))


def test_mark_application_valid_posting_upserts():
    with patch("export.export_tracker.query_rows", return_value=[{"1": 1}]):
        with patch("export.export_tracker.run_sql") as mock_run_sql:
            mark_application("adzuna-1", date(2026, 8, 1))
    sql, params = mock_run_sql.call_args.args
    assert "MERGE INTO applications" in sql
    assert params["posting_id"] == "adzuna-1"
    assert params["date_applied"] == "2026-08-01"
    assert params["status"] == "APPLIED"


def test_mark_application_invalid_status_raises_error():
    with pytest.raises(InvalidStatusError):
        mark_application("adzuna-1", date(2026, 8, 1), status="INVALID_STATUS")


def test_mark_application_with_custom_status_upserts():
    with patch("export.export_tracker.query_rows", return_value=[{"1": 1}]):
        with patch("export.export_tracker.run_sql") as mock_run_sql:
            mark_application("adzuna-1", date(2026, 8, 1), status="INTERVIEW")
    sql, params = mock_run_sql.call_args.args
    assert params["status"] == "INTERVIEW"
